import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))
import json
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
import fitz

from modules.detection.service import DetectionService


# =============================================================================
# CONFIGURATION
# =============================================================================

DOCUMENTS_DIR = Path("evaluation/documents")
OUTPUT_DIR = Path("evaluation/outputs")

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

service = DetectionService()


# =============================================================================
# ENTITY CATEGORIES
# =============================================================================

PII_TYPES = {
    "PERSON",
    "EMAIL",
    "PHONE_NUMBER",
    "ADDRESS",
    "AADHAAR_NUMBER",
    "PAN_NUMBER",
    "PASSPORT_NUMBER",
    "DRIVING_LICENSE",
    "BANK_ACCOUNT",
    "CREDIT_CARD",
    "IFSC_CODE",
    "UPI_ID",
    "PIN_CODE",
    "IP_ADDRESS",
    "URL",
    "SSN",
    "US_PHONE_NUMBER",
    "DATE_OF_BIRTH",
    "ZIP_CODE",
}

PHI_TYPES = {
    "PATIENT",
    "DOCTOR",
    "HOSPITAL",
    "MRN",
    "PROBLEM",
    "DISEASE",
    "CONDITION",
    "MEDICATION",
    "MEDICINE",
    "DRUG",
    "PROCEDURE",
    "SYMPTOM",
    "LAB_RESULT",
    "TEST",
    "MRN",
    "CLAIM_NUMBER",
    "INSURANCE_ID",
    "CPT_CODE",
    "ICD10_CODE",
}


# =============================================================================
# MAIN
# =============================================================================

def process_document(pdf_path: Path):

    print("=" * 90)
    print(f"Processing : {pdf_path.name}")
    print("=" * 90)

    pdf = fitz.open(pdf_path)

    report = {
        "document": pdf_path.name,
        "total_pages": len(pdf),
        "summary": {},
        "pages": [],
        "entities": [],
    }

    total_entities = 0

    total_pii = 0
    total_phi = 0

    pii_summary = defaultdict(int)
    phi_summary = defaultdict(int)

    detector_summary = defaultdict(int)

    # -------------------------------------------------------------------------

    for page_index in range(len(pdf)):

        page_number = page_index + 1

        page = pdf.load_page(page_index)

        text = page.get_text()

        if not text.strip():

            report["pages"].append(
                {
                    "page": page_number,
                    "pii": {},
                    "phi": {},
                }
            )

            continue

        print(f"\nProcessing Page {page_number}")

        entities = service.detect(
            text=text,
            page_number=page_number,
        )

        total_entities += len(entities)

        page_pii = defaultdict(list)
        page_phi = defaultdict(list)

        # -----------------------------------------------------------------

        for entity in entities:

            detector_summary[entity.detector] += 1

            item = {
                "page": entity.page_number,
                "entity_type": entity.entity_type,
                "entity_value": entity.entity_value,
                "confidence": round(entity.confidence_score, 3),
                "detector": entity.detector,
                "metadata": entity.metadata,
            }

            report["entities"].append(item)

            if entity.entity_type in PII_TYPES:

                total_pii += 1

                pii_summary[entity.entity_type] += 1

                page_pii[entity.entity_type].append(
                    entity.entity_value
                )

            elif entity.entity_type in PHI_TYPES:

                total_phi += 1

                phi_summary[entity.entity_type] += 1

                page_phi[entity.entity_type].append(
                    entity.entity_value
                )

        report["pages"].append(
            {
                "page": page_number,
                "pii": dict(page_pii),
                "phi": dict(page_phi),
            }
        )

    # -------------------------------------------------------------------------

    report["summary"] = {

        "total_entities": total_entities,

        "pii": {

            "total": total_pii,

            "breakdown": dict(sorted(pii_summary.items()))

        },

        "phi": {

            "total": total_phi,

            "breakdown": dict(sorted(phi_summary.items()))

        },

        "detectors": dict(sorted(detector_summary.items()))

    }

    return report


# =============================================================================

# =============================================================================
# PRINT REPORT
# =============================================================================

def print_report(report):

    print("\n")
    print("=" * 90)
    print("DOCUMENT INTELLIGENCE REPORT")
    print("=" * 90)

    print(f"Document        : {report['document']}")
    print(f"Pages           : {report['total_pages']}")
    print(f"Total Entities  : {report['summary']['total_entities']}")

    print()

    print("=" * 90)
    print("PII SUMMARY")
    print("=" * 90)

    print(f"Total PII : {report['summary']['pii']['total']}")
    print()

    for entity, count in report["summary"]["pii"]["breakdown"].items():
        print(f"{entity:<25} {count}")

    print()

    print("=" * 90)
    print("PHI SUMMARY")
    print("=" * 90)

    print(f"Total PHI : {report['summary']['phi']['total']}")
    print()

    for entity, count in report["summary"]["phi"]["breakdown"].items():
        print(f"{entity:<25} {count}")

    print()

    print("=" * 90)
    print("DETECTOR SUMMARY")
    print("=" * 90)

    for detector, count in report["summary"]["detectors"].items():
        print(f"{detector:<25} {count}")

    print()

    print("=" * 90)
    print("PAGE WISE REPORT")
    print("=" * 90)

    for page in report["pages"]:

        print()
        print("-" * 90)
        print(f"PAGE {page['page']}")
        print("-" * 90)

        #
        # PII
        #

        pii_total = sum(len(v) for v in page["pii"].values())

        print(f"\nPII ({pii_total})")

        if page["pii"]:

            for entity_type, values in page["pii"].items():

                print(f"\n{entity_type}")

                for value in values:

                    print(f"   • {value}")

        else:

            print("No PII detected.")

        #
        # PHI
        #

        phi_total = sum(len(v) for v in page["phi"].values())

        print(f"\nPHI ({phi_total})")

        if page["phi"]:

            for entity_type, values in page["phi"].items():

                print(f"\n{entity_type}")

                for value in values:

                    print(f"   • {value}")

        else:

            print("No PHI detected.")

    print()
    print("=" * 90)
# =============================================================================
# SAVE JSON
# =============================================================================

def save_json(report):

    output = OUTPUT_DIR / f"{Path(report['document']).stem}.json"

    with open(output, "w", encoding="utf-8") as f:

        json.dump(report, f, indent=4, ensure_ascii=False)

    print(f"\nJSON Saved : {output}")


# =============================================================================
# SAVE EXCEL
# =============================================================================

def save_excel(report):

    wb = Workbook()

    # ----------------------------------------------------------------------
    # SUMMARY
    # ----------------------------------------------------------------------

    ws = wb.active
    ws.title = "Summary"

    ws["A1"] = "Document Intelligence Report"
    ws["A1"].font = Font(bold=True)

    ws.append([])
    ws.append(["Document", report["document"]])
    ws.append(["Pages", report["total_pages"]])
    ws.append(["Total Entities", report["summary"]["total_entities"]])
    ws.append(["Total PII", report["summary"]["pii"]["total"]])
    ws.append(["Total PHI", report["summary"]["phi"]["total"]])

    # ----------------------------------------------------------------------
    # PII SUMMARY
    # ----------------------------------------------------------------------

    ws2 = wb.create_sheet("PII Summary")

    ws2.append(["Entity Type", "Count"])

    for cell in ws2[1]:
        cell.font = Font(bold=True)

    for entity, count in report["summary"]["pii"]["breakdown"].items():

        ws2.append([entity, count])

    # ----------------------------------------------------------------------
    # PHI SUMMARY
    # ----------------------------------------------------------------------

    ws3 = wb.create_sheet("PHI Summary")

    ws3.append(["Entity Type", "Count"])

    for cell in ws3[1]:
        cell.font = Font(bold=True)

    for entity, count in report["summary"]["phi"]["breakdown"].items():

        ws3.append([entity, count])

    # ----------------------------------------------------------------------
    # DETECTOR SUMMARY
    # ----------------------------------------------------------------------

    ws4 = wb.create_sheet("Detector Summary")

    ws4.append(["Detector", "Entities"])

    for cell in ws4[1]:
        cell.font = Font(bold=True)

    for detector, count in report["summary"]["detectors"].items():

        ws4.append([detector, count])

    # ----------------------------------------------------------------------
    # PAGE REPORT
    # ----------------------------------------------------------------------

    ws5 = wb.create_sheet("Page Report")

    ws5.append(
        [
            "Page",
            "Category",
            "Entity Type",
            "Detected Value",
        ]
    )

    for cell in ws5[1]:
        cell.font = Font(bold=True)

    for page in report["pages"]:

        #
        # PII
        #

        for entity_type, values in page["pii"].items():

            for value in values:

                ws5.append(
                    [
                        page["page"],
                        "PII",
                        entity_type,
                        value,
                    ]
                )

        #
        # PHI
        #

        for entity_type, values in page["phi"].items():

            for value in values:

                ws5.append(
                    [
                        page["page"],
                        "PHI",
                        entity_type,
                        value,
                    ]
                )

    output = OUTPUT_DIR / f"{Path(report['document']).stem}.xlsx"

    wb.save(output)

    print(f"Excel Saved : {output}")
if __name__ == "__main__":

    pdfs = list(DOCUMENTS_DIR.glob("*.pdf"))

    if not pdfs:

        print("No PDF files found.")
        exit()

    reports = []

    for pdf in pdfs:

        report = process_document(pdf)

        reports.append(report)

        print_report(report)

        save_json(report)

        save_excel(report)