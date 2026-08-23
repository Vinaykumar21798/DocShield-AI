from dataclasses import dataclass
from pathlib import Path
from time import perf_counter
from typing import Any, Dict, Optional

from database.models import Document


class NativeExtractionError(Exception):
    """
    Raised when native text extraction fails.
    """


@dataclass(frozen=True)
class TextExtractionResult:
    extracted_text: str
    page_count: int
    confidence_score: float
    processing_time: float
    extracted_text_path: Optional[str] = None
    structured_output: Optional[Dict[str, Any]] = None


class NativePDFExtractionService:
    """
    Extracts text from searchable PDF documents.
    """

    EMPTY_TEXT = ""
    SEARCHABLE_CONFIDENCE = 1.0
    EMPTY_CONFIDENCE = 0.0

    def extract(self, document: Document) -> TextExtractionResult:
        file_path = self._get_file_path(document)
        start_time = perf_counter()

        try:
            import fitz
        except ImportError as exc:
            raise NativeExtractionError(
                "PyMuPDF is required for native PDF extraction."
            ) from exc

        try:
            with fitz.open(file_path) as pdf_document:
                page_count = pdf_document.page_count
                extracted_text = self._extract_document_text(pdf_document)

        except Exception as exc:
            raise NativeExtractionError(
                f"Failed to extract native PDF text: {file_path}"
            ) from exc

        processing_time = perf_counter() - start_time
        confidence_score = self._get_confidence_score(extracted_text)

        return TextExtractionResult(
            extracted_text=extracted_text,
            page_count=page_count,
            confidence_score=confidence_score,
            processing_time=processing_time,
        )

    def _extract_document_text(self, pdf_document) -> str:
        page_texts = []

        for page_number in range(pdf_document.page_count):
            page = pdf_document.load_page(page_number)
            page_text = self._extract_page_layout_text(page)

            if page_text:
                page_texts.append(page_text)

        return "\n\n\f\n\n".join(page_texts).strip()

    @staticmethod
    def _extract_page_layout_text(page) -> str:
        try:
            tabs = page.find_tables()
            if not tabs or not tabs.tables:
                return page.get_text("text").strip()

            import fitz
            table_entries = []
            table_rects = []
            for tab in tabs:
                t_rect = fitz.Rect(tab.bbox)
                table_rects.append(t_rect)
                try:
                    md = tab.to_markdown().strip()
                except Exception:
                    extracted = tab.extract()
                    md = "\n".join(["\t".join(str(c or "") for c in row) for row in extracted])
                table_entries.append((t_rect.y0, md))

            blocks = page.get_text("blocks")
            elements = []
            for b in blocks:
                b_rect = fitz.Rect(b[:4])
                b_text = b[4].strip()
                if not b_text:
                    continue
                inside_table = False
                for t_rect in table_rects:
                    intersect = b_rect & t_rect
                    if intersect.get_area() > 0.5 * b_rect.get_area():
                        inside_table = True
                        break
                if not inside_table:
                    elements.append((b_rect.y0, b_text))

            for y0, md in table_entries:
                elements.append((y0, md))

            elements.sort(key=lambda item: item[0])
            return "\n\n".join(item[1] for item in elements).strip()
        except Exception:
            return page.get_text("text").strip()

    def _get_file_path(self, document: Document) -> Path:
        if not document.storage_path:
            raise NativeExtractionError(
                f"Document storage path is missing: {document.id}"
            )

        file_path = Path(document.storage_path)

        if not file_path.exists():
            raise NativeExtractionError(
                f"Stored document file not found: {file_path}"
            )

        return file_path

    def _get_confidence_score(self, extracted_text: str) -> float:
        if extracted_text.strip():
            return self.SEARCHABLE_CONFIDENCE

        return self.EMPTY_CONFIDENCE


class NativeTextExtractionService:
    """
    Extracts text from text-based document formats.
    """

    TEXT_EXTENSIONS = {".txt"}
    DOCX_EXTENSIONS = {".docx"}
    SEARCHABLE_CONFIDENCE = 1.0
    EMPTY_CONFIDENCE = 0.0

    def extract(self, document: Document) -> TextExtractionResult:
        file_path = self._get_file_path(document)
        extension = self._get_extension(document, file_path)
        start_time = perf_counter()

        if extension in self.TEXT_EXTENSIONS:
            extracted_text = self._extract_plain_text(file_path)

        elif extension in self.DOCX_EXTENSIONS:
            extracted_text = self._extract_docx_text(file_path)

        else:
            raise NativeExtractionError(
                f"Unsupported native text document type: {extension}"
            )

        processing_time = perf_counter() - start_time

        return TextExtractionResult(
            extracted_text=extracted_text,
            page_count=1,
            confidence_score=self._get_confidence_score(extracted_text),
            processing_time=processing_time,
        )

    @staticmethod
    def _extract_plain_text(file_path: Path) -> str:
        return file_path.read_text(encoding="utf-8", errors="replace").strip()

    @staticmethod
    def _extract_docx_text(file_path: Path) -> str:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise NativeExtractionError(
                "python-docx is required for DOCX text extraction."
            ) from exc

        try:
            docx_document = DocxDocument(str(file_path))
        except Exception as exc:
            raise NativeExtractionError(
                f"Failed to extract DOCX text: {file_path}"
            ) from exc

        return "\n".join(
            paragraph.text.strip()
            for paragraph in docx_document.paragraphs
            if paragraph.text.strip()
        ).strip()

    def _get_file_path(self, document: Document) -> Path:
        if not document.storage_path:
            raise NativeExtractionError(
                f"Document storage path is missing: {document.id}"
            )

        file_path = Path(document.storage_path)

        if not file_path.exists():
            raise NativeExtractionError(
                f"Stored document file not found: {file_path}"
            )

        return file_path

    def _get_extension(
        self,
        document: Document,
        file_path: Path,
    ) -> str:
        candidates = (
            file_path.suffix,
            Path(document.stored_filename or "").suffix,
            Path(document.filename or "").suffix,
        )

        for extension in candidates:
            normalized_extension = extension.lower().strip()
            if normalized_extension:
                return normalized_extension

        return ""

    def _get_confidence_score(self, extracted_text: str) -> float:
        if extracted_text.strip():
            return self.SEARCHABLE_CONFIDENCE

        return self.EMPTY_CONFIDENCE